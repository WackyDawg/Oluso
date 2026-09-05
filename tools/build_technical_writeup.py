from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = "102A43"
DEEP = "07182B"
TEAL = "0AAE9B"
CYAN = "239CC4"
PALE = "EAF4F7"
LIGHT = "F4F7FA"
AMBER = "E39A19"
RED = "C44949"
WHITE = "FFFFFF"
MUTED = "52697D"


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    fill = tc_pr.find(qn("w:shd"))
    if fill is None:
        fill = OxmlElement("w:shd")
        tc_pr.append(fill)
    fill.set(qn("w:fill"), color)


def set_cell_border(cell, color: str = "C7D5E0", size: str = "5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def keep(paragraph) -> None:
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = True


def add_text(paragraph, text: str, *, bold: bool = False, color: str | None = None,
             size: float | None = None) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)


def body(doc: Document, text: str, *, space_after: float = 4.0) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.style = doc.styles["Body Text"]
    paragraph.paragraph_format.space_after = Pt(space_after)


def bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.45)
    paragraph.paragraph_format.first_line_indent = Cm(-0.25)
    paragraph.paragraph_format.space_after = Pt(2.0)
    add_text(paragraph, text)


def heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(5 if level == 1 else 3)
    paragraph.paragraph_format.space_after = Pt(3)
    keep(paragraph)


def metric_row(doc: Document, metrics: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=len(metrics))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for cell, (label, value, detail) in zip(table.rows[0].cells, metrics):
        cell.width = Cm(6.1)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, NAVY)
        set_cell_border(cell, NAVY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(1)
        add_text(paragraph, label.upper() + "\n", bold=True, color="B7D1E3", size=7.5)
        add_text(paragraph, value + "\n", bold=True, color=WHITE, size=16)
        add_text(paragraph, detail, color="DCE9F2", size=7.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def compact_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float],
                  first_col_bold: bool = True) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, (cell, label) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Cm(widths[index])
        shade(cell, NAVY)
        set_cell_border(cell, NAVY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, label, bold=True, color=WHITE, size=7.5)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        tr_pr = cells[0]._tc.getparent().get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for index, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Cm(widths[index])
            shade(cell, LIGHT if row_index % 2 == 0 else WHITE)
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, value, bold=first_col_bold and index == 0, color=DEEP, size=7.25)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def banner(doc: Document, text: str, color: str = TEAL) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, PALE)
    set_cell_border(cell, color, "10")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_text(p, text, bold=True, color=NAVY, size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def masthead(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.cell(0, 0).width = Cm(12.8)
    table.cell(0, 1).width = Cm(5.0)
    for cell in table.rows[0].cells:
        shade(cell, NAVY)
        set_cell_border(cell, NAVY)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    add_text(left, "ICSC UNIVERSITIES CATEGORY · TRACK A\n", bold=True, color=TEAL, size=8)
    add_text(left, "AEGISTWIN", bold=True, color=WHITE, size=19)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    add_text(right, "TECHNICAL WRITE-UP\n", bold=True, color="C5D8E6", size=8)
    add_text(right, "v1.4 · 30 AUG 2026", bold=True, color=WHITE, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_title(doc: Document, number: str, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    add_text(p, number + "  ", bold=True, color=TEAL, size=10)
    add_text(p, title, bold=True, color=NAVY, size=19)
    p = doc.add_paragraph(subtitle)
    p.paragraph_format.space_after = Pt(5)
    p.runs[0].font.size = Pt(8.8)
    p.runs[0].font.color.rgb = RGBColor.from_string(MUTED)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    section.header_distance = Cm(0.55)
    section.footer_distance = Cm(0.55)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(8.3)
    normal.font.color.rgb = RGBColor.from_string(DEEP)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 0.99
    body_style = styles["Body Text"]
    body_style.font.name = "Aptos"
    body_style.font.size = Pt(8.3)
    body_style.font.color.rgb = RGBColor.from_string(DEEP)
    body_style.paragraph_format.line_spacing = 0.99
    for name, size, color in (("Title", 27, NAVY), ("Heading 1", 12, NAVY), ("Heading 2", 9.5, TEAL)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(
        footer,
        "AegisTwin · Track A · Synthetic prototype · Technical submission",
        color=MUTED,
        size=7,
    )


def build(project: Path, output: Path) -> None:
    evaluation = json.loads((project / "artifacts/evaluation.json").read_text())
    latency = json.loads((project / "artifacts/latency.json").read_text())
    adaptive = json.loads((project / "artifacts/adaptive_redteam.json").read_text())
    v5_demo = json.loads((project / "artifacts/v5_feature_demo.json").read_text())
    load = json.loads((project / "artifacts/concurrent_load.json").read_text())
    outage = json.loads((project / "artifacts/outage_resilience.json").read_text())
    agent = json.loads((project / "artifacts/agent_terminal_demo.json").read_text())
    sketch = json.loads((project / "artifacts/fraud_sketch_exchange.json").read_text())
    robustness = json.loads((project / "artifacts/robustness.json").read_text())
    challenge = evaluation["operating_points"]["Customer confirmation or stronger"]
    scenarios = evaluation["scenario_cohorts"]
    hard = evaluation["hard_negative_intervention"]
    model_hash = hashlib.sha256((project / "models/ato_model.joblib").read_bytes()).hexdigest()
    data_hash = hashlib.sha256((project / "data/training_features.csv").read_bytes()).hexdigest()

    doc = Document()
    configure(doc)
    props = doc.core_properties
    props.title = "AegisTwin Track A Technical Write-up"
    props.subject = "ICSC Universities Category — Spotting Account Takeover From Behaviour"
    props.keywords = "ICSC, account takeover, behavioural AI, USSD, Nigeria, privacy, fraud sketch"
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""

    masthead(doc)
    title = doc.add_paragraph()
    title.style = doc.styles["Title"]
    title.paragraph_format.space_after = Pt(2)
    add_text(title, "Spotting account takeover by recognising the human", bold=True, color=NAVY, size=24)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_text(p, "A working, explainable defence for Nigerian app, USSD and agency banking", bold=True, color=TEAL, size=10.5)

    metric_row(doc, [
        ("Held-out recall", f"{challenge['recall']:.1%}", f"{challenge['true_positives']}/{challenge['takeover_events']} challenged"),
        ("False-positive rate", f"{challenge['false_positive_rate']:.3%}", f"{challenge['false_positives']}/{challenge['normal_events']} normal"),
        ("End-to-end p99", f"{latency['p99_ms']:.0f} ms", f"{latency['requests']} warm full-API calls"),
    ])

    heading(doc, "Executive summary")
    body(doc, "PINs and OTPs prove knowledge, not presence. AegisTwin maintains a server-side behavioural twin for each account, compares a live transaction with that person's established habits, and fuses the result with a calibrated population model. It then chooses the least disruptive reversible response and explains it in one plain sentence.")
    banner(doc, f"Core claim: at the challenge threshold, AegisTwin caught {challenge['true_positives']} of {challenge['takeover_events']} held-out takeovers while challenging {challenge['false_positives']} of {challenge['normal_events']:,} normal sessions.")

    heading(doc, "What makes the prototype differentiated")
    bullet(doc, "Two-sided network intelligence: a Mule Graph catches many-sender/rapid-cash-out behaviour; analyst-confirmed recipient reputation propagates across accounts with decay and a two-victim hard-floor rule.")
    bullet(doc, "Agent-Terminal Integrity Twin: gateway-attested cross-customer patterns expose one compromised agency endpoint, while a compound guard protects legitimate high-throughput agents.")
    bullet(doc, "Decision Confidence Envelope: risk and certainty are separate; thin history, missing evidence and scorer disagreement are surfaced instead of disguised as a safe score.")
    bullet(doc, "Cross-Channel + Calendar Twins: unusual app-to-USSD enrolment sequences raise takeover risk, while three-cycle monthly payments reduce needless friction.")
    bullet(doc, "Coercion + trajectory defence: consented in-call safety signals can privately pause payment; recovery/auth/SIM precursors create a decaying account risk window.")
    bullet(doc, "Personalised friction and recourse: policy compares fraud exposure with customer cost, respects a Regret Budget, and returns safe, non-gameable clearance routes.")
    bullet(doc, "Adaptive attacker: 2,688 valid raw-event variants expose a NGN 48,000 evasion and test recipient/precursor mitigations without self-training on the answer.")
    bullet(doc, "AegisMesh + outage assurance: signed rotating cross-bank sketches share patterns, not records; local-action ceilings, revocation, signed cache, no-settlement receipt and hash-chained recovery constrain harm.")

    heading(doc, "Challenge fit")
    compact_table(doc, ["Required element", "Delivered"], [
        ["Synthetic sessions", "30,000 scored sessions, 750 personas, eight attack families, 12 hard-legitimate families, delayed feedback and causal precursors."],
        ["Live model", "FastAPI service; server-built features; calibrated Random Forest + transparent personal anomaly scorer + policy fusion."],
        ["Both error sides", "Recall, precision, false-positive rate, TP/FP/FN, cohort results, rare-prevalence projection and known failures."],
        ["Plain explanation", "One customer-ready sentence plus safe channel, clearance action and expected timing for every interruption."],
        ["USSD / feature phone", "Gateway timing/menu path, bank history, recipient graph and optional attested telco signals; no app sensor dependency."],
        ["Power / network cuts", "Working four-state outage controller; degraded confidence, isolated no-settlement handling, tamper detection and idempotent reconciliation."],
        ["Agency / consortium", "Attested terminal risk plus signed private recipient/terminal/campaign sketches; one bank or uncorroborated evidence only monitors."],
    ], [4.0, 13.6])

    doc.add_page_break()
    add_page_title(doc, "02", "System and AI design", "A hybrid decision is stronger than either a black-box model or a fixed rule list alone.")
    compact_table(doc, ["1 · Event", "2 · Sender twin", "3 · Other risk views", "4 · AI + policy", "5 · Evidence"], [[
        "App / USSD / agent",
        "Sender behavioural twin",
        "Graph + agent + AegisMesh",
        "Hybrid AI + cost policy",
        "Recourse + audit hash",
    ]], [3.45, 3.45, 3.55, 3.55, 3.55], first_col_bold=False)

    heading(doc, "Signals and trust boundaries")
    compact_table(doc, ["Signal family", "Examples", "How trust is constrained"], [
        ["Transaction", "amount/MAD, balance drain, recipient novelty, velocity, hour", "Computed inside the bank from earlier events; client cannot supply its baseline."],
        ["App behaviour", "typing, handling, navigation, coercion-safety telemetry", "Optional; coercion fields are zeroed unless consented and device-attested."],
        ["USSD behaviour", "menu depth, interaction speed, path signature, SIM/IP", "Collected by gateway/core systems, usable on feature phones."],
        ["Telco lifecycle", "activation age, IMSI/ICCID change flags, swaps/30d, OTP proximity", "Ignored unless gateway-attested; accepts bounded indicators, not raw subscriber IDs."],
        ["Account + recipient", "failed auth, recovery, mule fan-in/cash-out, verified change", "Precursor hazard decays; mule floor requires cash-out, not merchant fan-in alone."],
        ["Agent terminal", "customer diversity, recipient concentration, auth failures, location/shift", "Ignored unless gateway-attested; one victim only monitors, three independent victims quarantine."],
        ["Exchange + policy", "rotating sketch, issuer count, confidence, action costs", "Signed banks share no raw record; without local corroboration the exchange can only monitor."],
    ], [3.1, 6.9, 7.6])

    heading(doc, "Decision mechanics")
    body(doc, "The population model learns attack combinations; personal, recipient and agent-terminal scorers expose accountable deviations. Ordered channel transitions, calendar rhythm and a decaying precursor window add sequence and time. AegisMesh adds delayed cross-bank evidence after the model, never as a leaked training label. Risk fusion is followed by separate confidence and a friction-cost policy. Shared evidence alone can monitor; locally corroborated evidence can contribute a delay but not a hold.")
    banner(doc, f"Adaptive result: NGN {adaptive['best_evasion']['config']['amount']:,.0f} evasion scores {adaptive['best_evasion']['outcome']['score']:.3f}; recipient graph raises the same event to {adaptive['recipient_graph_mitigation']['score']:.3f} and a reversible delay.", CYAN)
    body(doc, f"Live proof: cross-channel sequence {v5_demo['cross_channel_sequence']['risk_score']:.3f}; a three-cycle payment remains {v5_demo['calendar_twin']['risk_score']:.3f}. AegisMesh moves a baseline {sketch['baseline']['risk_score']:.3f} to one-bank monitoring {sketch['one_institution']['other_account_score']:.3f}; two uncorroborated banks cap at {sketch['two_institutions_without_local_corroboration']['risk_score']:.3f}, while local corroboration selects a reversible delay at {sketch['two_institutions_with_local_corroboration']['risk_score']:.3f}.", space_after=2)
    body(doc, f"Outage proof: three events were hash-chained; deliberate tampering was detected; all {outage['reconciliation']['processed']} events re-scored after recovery with {outage['reconciliation']['duplicate_actions']} duplicate actions and {outage['reconciliation']['automatic_settlements']} automatic settlements. Detection {outage['objectives']['outage_detection_ms']:.1f} ms; recovery {outage['objectives']['recovery_and_reconciliation_ms']:.1f} ms on the test laptop.", space_after=2)
    body(doc, f"Agent proof: a busy legitimate terminal remains at {agent['busy_legitimate_terminal']['risk_score']:.3f}; a compromised cross-customer terminal reaches {agent['compromised_terminal']['risk_score']:.3f} and a reversible delay. Three independent victims quarantine it; an unattested forged claim contributes {agent['forged_unattested_claim']['agent_evidence_coverage']:.0%} agent evidence.", space_after=2)

    heading(doc, "Why it remains explainable")
    bullet(doc, "The API returns model/anomaly/fused risk, a separate decision-confidence score and reasons, profile/evidence coverage, recipient status, action costs and uncertainty note.")
    bullet(doc, "Customer wording avoids hidden technical claims; structured recourse offers registered-device confirmation, official callback or safe expiry without disclosing thresholds.")
    bullet(doc, "If the model artifact is missing or has a mismatched feature schema, runtime falls back to transparent scoring rather than crashing or accepting blindly.")

    heading(doc, "Operational flow")
    compact_table(doc, ["Risk", "Customer treatment", "Safeguard"], [
        ["Low", "Allow", "Current event waits in learning quarantine"],
        ["Guarded", "Allow + monitor", "No payment interruption"],
        ["Elevated", "Trusted-channel confirmation", "Do not use the possibly compromised channel"],
        ["High", "15-minute settlement delay", "Customer can confirm/cancel; reversible"],
        ["Critical", "Two-hour hold + review", "Analyst review; SHA-256 chained evidence"],
        ["Coercion", "Private 15-minute safety pause", "End calls/screen share; confirm through an official channel"],
        ["Outage", "Monitor / confirm / pending", "Never treat missing evidence as safe; never claim isolated settlement"],
    ], [2.8, 6.0, 8.8])

    doc.add_page_break()
    add_page_title(doc, "03", "Synthetic data, evaluation and failure cases", "The test is chronological and the numbers describe the exact final fusion policy.")
    heading(doc, "How the dataset was generated")
    body(doc, "A deterministic generator (seed 2026) creates 750 personas across six Nigerian cities. After eight earlier events, accounts may receive one of eight attacks or 12 hard legitimate changes. It includes groomed-recipient takeovers plus legitimate recovery, failed-login retry, cooperative networks, mule cash-out, shared agents, monthly cycles and delayed confirmations. Post-label reputation and event-type shortcuts are excluded from population-model training. The newest 18% is held out.")

    heading(doc, "Final held-out results")
    compact_table(doc, ["Live threshold", "Recall", "Precision", "FPR", "TP / FP / FN"], [
        [name, f"{item['recall']:.1%}", f"{item['precision']:.1%}", f"{item['false_positive_rate']:.3%}", f"{item['true_positives']} / {item['false_positives']} / {item['false_negatives']}"]
        for name, item in evaluation["operating_points"].items()
    ], [6.1, 2.4, 2.4, 2.5, 4.2])
    body(doc, f"Held-out window: {evaluation['test_samples']:,} sessions, {evaluation['takeover_samples']} takeovers ({evaluation['synthetic_prevalence']:.2%}); model/fused ROC-AUC {evaluation['model_only_roc_auc']:.4f}/{evaluation['fused_roc_auc']:.4f}, neither perfect. The best single feature is {robustness['chronological_separability']['top_univariate_features'][0]['separation_auc']:.4f}. At 0.1% prevalence, the measured confirmation rates project {evaluation['prevalence_stress']['0.100%']['projected_false_alerts_per_10000']:.2f} false alerts per 10,000; this is a mathematical projection, not field validation.")

    heading(doc, "Channel and scenario evidence")
    compact_table(doc, ["Cohort", "Held-out evidence", "Result at ≥0.48"], [
        ["USSD", f"{evaluation['channel_cohorts']['ussd']['samples']:,} sessions / {evaluation['channel_cohorts']['ussd']['takeovers']} takeovers", f"{evaluation['channel_cohorts']['ussd']['recall']:.0%} recall; {evaluation['channel_cohorts']['ussd']['false_positive_rate']:.3%} FPR"],
        ["App", f"{evaluation['channel_cohorts']['app']['samples']:,} sessions / {evaluation['channel_cohorts']['app']['takeovers']} takeovers", f"{evaluation['channel_cohorts']['app']['recall']:.0%} recall; {evaluation['channel_cohorts']['app']['false_positive_rate']:.3%} FPR"],
        ["Agent", f"{evaluation['channel_cohorts']['agent']['samples']:,} sessions / {evaluation['channel_cohorts']['agent']['takeovers']} takeovers", f"{evaluation['channel_cohorts']['agent']['recall']:.0%} recall; {evaluation['channel_cohorts']['agent']['false_positive_rate']:.3%} FPR"],
        ["Unseen accounts", f"{robustness['account_disjoint']['test_accounts']} accounts / {robustness['account_disjoint']['takeover_events']} takeovers", f"{robustness['account_disjoint']['customer_confirmation']['recall']:.0%} recall; {robustness['account_disjoint']['customer_confirmation']['false_positive_rate']:.3%} FPR"],
        ["Busy agent", f"{hard['busy_agent_merchant_payment']['events']} legitimate events", f"{hard['busy_agent_merchant_payment']['intervention_rate']:.0%} challenged"],
        ["Community cooperative", f"{hard['community_collection_payment']['events']} legitimate events", f"{hard['community_collection_payment']['intervention_rate']:.0%} challenged"],
    ], [4.2, 6.5, 6.9])

    heading(doc, "Known failures—shown, not hidden")
    bullet(doc, f"Low-and-slow recall is {scenarios['low_and_slow']['recall']:.0%} ({round(scenarios['low_and_slow']['takeovers'] * scenarios['low_and_slow']['recall'])}/{scenarios['low_and_slow']['takeovers']}); remote-control USSD recall is {scenarios['remote_control_ussd']['recall']:.0%}. These deliberately subtle misses replace the previous unrealistic perfect ranking.")
    bullet(doc, f"Account-disjoint confirmation recall falls to {robustness['account_disjoint']['customer_confirmation']['recall']:.1%} with {robustness['account_disjoint']['account_overlap']} customer overlap, showing the cost of generalising beyond established twins.")
    bullet(doc, f"Emergency transfers are the principal hard-negative cohort: {hard['emergency_transfer']['intervention_rate']:.2%} ({round(hard['emergency_transfer']['events'] * hard['emergency_transfer']['intervention_rate'])}/{hard['emergency_transfer']['events']}) were challenged. Every action remains reversible.")
    bullet(doc, f"Adaptive search found a same-device/SIM/IP NGN 48,000 evasion at {adaptive['best_evasion']['outcome']['score']:.3f}. A fresh recipient without graph or precursor evidence remains a disclosed weakness.")

    doc.add_page_break()
    add_page_title(doc, "04", "Security, deployment and reproducibility", "A prototype judges can run today, with explicit gates before any real-money use.")
    heading(doc, "Security and privacy controls")
    compact_table(doc, ["Risk", "Implemented control", "Production next step"], [
        ["Forged client baselines", "All behavioural profiles computed server-side", "Authenticated event bus + schema registry"],
        ["Forged telco claims", "Unattested envelope is zeroed", "mTLS, signed claims, issuer allow-list, key rotation"],
        ["Forged / accused agent terminal", "Unattested evidence ignored; 1 victim monitors, 3 quarantine", "Gateway signatures, appeal/reinstatement SLA and cohort calibration"],
        ["Poisoning", "24h quarantine; fraud revokes event and rebuilds profile", "Durable scheduler and labelled shadow-mode review"],
        ["Bad feedback", "Independent second approval; appeals; 90-day decay", "Identity provider, reason-quality QA and appeal SLA"],
        ["Replay / tampering", "Exact causal replay + idempotent lifecycle + signed audit head", "WORM anchor, durable queue and HSM keys"],
        ["Private exchange abuse", "Per-bank signature, nonce, rotation, rate, expiry, revocation, action ceiling", "VOPRF/PSI, HSM keys and consortium appeal governance"],
        ["Tenant / role abuse", "Tenant-bound queries; integration/analyst/auditor separation", "Workload identity, mTLS and central SIEM"],
        ["Privacy leakage", "Synthetic data; coarse/derived signals; no raw IMSI/ICCID", "Consent, minimization, retention and access review"],
        ["Customer harm", "Reversible actions + Regret Budget + plain explanation", "Bank-specific appeals and emergency handling"],
        ["Power / network loss", "Signed modes + edge capsule + confidence floor + chained journal", "mTLS/HSM keys, encrypted replicated log and alternate-site drill"],
        ["Release configuration", "JSON/CSV list parser + checked-in example + clean-package boot test", "Secrets manager, deployment admission checks and staged rollout"],
    ], [3.5, 6.7, 7.4])

    heading(doc, "Reproduce the evidence")
    compact_table(doc, ["Command", "Purpose"], [
        ["pip install -e '.[dev,dashboard]'", "Install local prototype"],
        ["python scripts/demo_fraud_sketch_exchange.py", "Prove signed private propagation, ceilings, outage cache, tamper and revocation"],
        ["pytest --cov=aegistwin && python scripts/release_smoke_test.py .", "Run 64 tests, coverage gate and clean-copy environment/API boot"],
        ["python scripts/generate_synthetic.py --rows 30000", "Regenerate raw sessions and derived features"],
        ["python scripts/train_model.py && python scripts/evaluate_system.py", "Retrain and reproduce held-out policy metrics"],
        ["python scripts/evaluate_robustness.py", "Reject perfect synthetic ranking; reproduce zero-overlap account test"],
        ["python scripts/adaptive_attack_search.py", "Reproduce the 2,688-candidate adversarial search"],
        ["python scripts/demo_outage_resilience.py && python scripts/demo_agent_terminal.py", "Prove outage recovery and agent-terminal safeguards"],
        ["python scripts/load_test.py --requests 100 --workers 8", f"Reproduce {load['successes']}/{load['requests']} concurrent successes; {load['latency_ms']['p99']:.0f} ms p99"],
    ], [7.6, 10.0])

    heading(doc, "Evidence provenance")
    body(doc, f"Model SHA-256: {model_hash[:24]}…  ·  Training-feature CSV SHA-256: {data_hash[:24]}…  ·  Model version: {evaluation['model_version']}. Full hashes and machine-readable results are included in the package.")

    heading(doc, "External idea review and clean-room boundary")
    body(doc, "The public sliit-msc-research SIM-swap repository suggested lifecycle categories but was not copied or imported: it had no licence or generator, its file naming and row counts conflicted, roughly 73.9% of rows were labelled fraud, and its README derived labels from the same inputs. AegisTwin independently implements bounded attested concepts and rare longitudinal evaluation; the full source audit is included.")

    heading(doc, "Submission assets and final gate")
    body(doc, "Included: four-page PDF/DOCX, captioned real-run demo, source ZIP, synthetic data, model, 64 tests, evaluation/robustness/latency/load/adaptive/platform/outage/agent/sketch JSON, CycloneDX SBOM, restore manifest and checksums. The captain must publish the source and replace CODE_LINK.md. Synthetic results must not be represented as bank-field accuracy.")
    banner(doc, "Recommended next pilot: shadow-mode scoring plus a two-bank governed sketch sandbox; calibrate customer cost and privacy leakage before any real-money action.", AMBER)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, default=Path.cwd())
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.project.resolve(), args.output.resolve())
    print(f"wrote={args.output.resolve()}")


if __name__ == "__main__":
    main()
